"""Run: pytest -q tests/test_local_backend.py"""
import io
import sys
from pathlib import Path
from unittest.mock import patch

import fitz
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
from main import app

client = TestClient(app)


def sample_pdf():
    pdf = fitz.open()
    page = pdf.new_page(width=240, height=150)
    page.insert_text((20, 50), "Hello PDF HDV", fontsize=15)
    data = pdf.tobytes()
    pdf.close()
    return data


def post(path, body=None, **fields):
    return client.post(path, files={"file": ("drawing.pdf", body or sample_pdf(), "application/pdf")}, data=fields)


def test_capabilities():
    result = client.get("/api/capabilities")
    assert result.status_code == 200
    assert set(["txt", "docx", "tiff"]).issubset(set(result.json()["convert"]))


def test_pdf_txt_unicode_header():
    result = post("/api/convert", target="txt")
    assert result.status_code == 200
    assert "Hello PDF HDV" in result.content.decode("utf-8-sig")
    assert result.headers["cache-control"] == "no-store"


def test_pdf_docx_editable():
    result = post("/api/convert", target="docx")
    assert result.status_code == 200, result.text
    word = Document(io.BytesIO(result.content))
    assert "Hello PDF HDV" in " ".join(p.text for p in word.paragraphs)


def test_pdf_tiff_raster():
    result = post("/api/convert", target="tiff")
    assert result.status_code == 200, result.text
    with Image.open(io.BytesIO(result.content)) as image:
        assert image.format == "TIFF" and image.width > 200


def test_compress_pdf_still_valid_and_no_bigger():
    sample = sample_pdf()
    result = post("/compress", body=sample, level="Mạnh")
    assert result.status_code == 200
    assert len(result.content) <= len(sample)
    with fitz.open(stream=result.content, filetype="pdf") as pdf:
        assert "Hello PDF HDV" in pdf[0].get_text()


def test_bad_data_and_target():
    assert post("/api/convert", body=b"BAD", target="txt").status_code == 415
    assert post("/api/convert", target="pdfa").status_code == 400
    assert post("/api/convert", body=b"BAD", target="pdf").status_code == 415


def test_ocr_language_is_checked_before_processing():
    with patch("main._ocr_languages", return_value=["eng"]), patch("main._require_tool", return_value="tesseract"):
        assert post("/api/ocr", target="txt", language="vie+eng").status_code == 422
    assert post("/api/ocr", target="txt", language="../tmp").status_code == 400


def test_upload_limit():
    result = post("/api/convert", body=b"a" * (50 * 1024 * 1024 + 1), target="txt")
    assert result.status_code == 413


def test_office_roundtrip():
    import shutil
    if not shutil.which("soffice"):
        return
    word = Document()
    word.add_paragraph("Sample office HDV")
    stream = io.BytesIO()
    word.save(stream)
    response = client.post("/api/convert", files={"file": ("sample.docx", stream.getvalue(),
                                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                                    data={"target": "pdf"})
    assert response.status_code == 200, response.text
    with fitz.open(stream=response.content, filetype="pdf") as pdf:
        assert "Sample office HDV" in pdf[0].get_text()
